import pandas as pd
from pathlib import Path
import docx
import pptx
import fitz
import openpyxl
from tqdm import tqdm
from datetime import datetime

from pipeline.taskfactory import TaskWithInputFileMonitor

import sys
sys.path.append('../src/general/')
from checkAuthorNames import NameChecker

class pdfMetaExtractor:
    def __init__(self, path, file_type):
        self.path = path

    def extract(self):
        try:
            document = fitz.open(self.path)
        except Exception as e:
            print(f"Error extracting metadata from {self.path}: {e}")
            return None
        metadata = document.metadata
        if metadata is None:
            return None
        modifided_date = self.get_creation_date(metadata['modDate'])
        if modifided_date is None:
            modifided_date = pd.NaT
        created_date = self.get_creation_date(metadata['creationDate'])
        if created_date is None:
            created_date = pd.NaT
        metadata_dict = {
            'file:author': metadata['author'],
            'file:keywords': metadata['keywords'],
            'file:subject': metadata['subject'],
            'file:title': metadata['title'],
            'file:created': created_date,
            'file:modified': modifided_date,
            #'file:creator': metadata['creator'],
            #'file:producer': metadata['producer'],
            #'file:format': metadata['format'],
        }
        return metadata_dict

    def get_creation_date(self, x):
        x = x.replace("'","")
        if x == "":
            return None
        if x[-1] == "Z":       #D:20210421114907Z
            x = x[:-1]
        if "D" in x:
            if len(x) == 21:
                if "Z" in x: 
                    x=x.split("Z")[0]
                    pattern ="D:%Y%m%d%H%M%S"
                if "+" in x or "-" in x:
                    pattern ="D:%Y%m%d%H%M%S%z"
            if len(x) == 16:   #D:20201104115242  
                pattern ="D:%Y%m%d%H%M%S"
        try:
            return datetime.strptime(x, pattern)
        except:
            return None


class officeMetaExtractor:
    def __init__(self, path, file_type):
        self.path = path
        self.valid_file=False
        if file_type == 'docx':
            try:
                self.document = docx.Document(self.path)
                self.valid_file=True
            except Exception as e:
                print(f"Error loading {self.path}: {e}")
        if file_type == 'pptx':
            try:
                self.document = pptx.Presentation(self.path)
                self.valid_file=True
            except Exception as e:
                print(f"Error loading {self.path}: {e}")

    def extract(self):
        if not self.valid_file:
            return None
        try:
            metadata = self.document.core_properties
        except Exception as e:
            print(f"Error extracting metadata from {self.path}: {e}")
            return None
        
        metadata_dict = { 
            'file:author': metadata.author,
            'file:keywords': metadata.keywords,
            'file:subject': metadata.subject,
            'file:title': metadata.title,
            'file:created': metadata.created,
            'file:modified': metadata.modified,
            #'file:last_modified_by': metadata.last_modified_by,
            #'file:category': metadata.category,
            #'file:content_status': metadata.content_status,
            'file:language': metadata.language,
            #'file:version': metadata.version,
            #'file:revision': metadata.revision,
        }
        return metadata_dict

class xlsxMetaExtractor:
    def __init__(self, path, file_type):
        self.path = path
        self.valid_file=False
        if file_type == 'xlsx':
            try:
                self.document = openpyxl.load_workbook(self.path)
                self.valid_file=True
            except Exception as e:
                print(f"Error loading {self.path}: {e}")

    def extract(self):
        if not self.valid_file:
            return None
        try:
            metadata = self.document.properties
        except Exception as e:
            print(f"Error extracting metadata from {self.path}: {e}")
            return None
        
        # <openpyxl.packaging.core.DocumentProperties object>
        # Parameters:
        # creator=u'User', title=None, description=None, subject=None, identifier=None,
        # language=None, created=datetime.datetime(2018, 12, 11, 9, 55, 2),
        # modified=datetime.datetime(2018, 12, 11, 10, 30, 38), lastModifiedBy=u'User',
        # category=None, contentStatus=None, version=None, revision=None, keywords=None,
        # lastPrinted=None

        metadata_dict = { 
            'file:author': metadata.creator,
            'file:keywords': metadata.keywords,
            'file:subject': metadata.subject,
            'file:title': metadata.title,
            'file:created': metadata.created,
            'file:modified': metadata.modified,
        }
        return metadata_dict

extractors = {
    'pptx': officeMetaExtractor,
    'docx': officeMetaExtractor,
    'pdf': pdfMetaExtractor,
    'xlsx': xlsxMetaExtractor,
}

class MetaDataExtraction(TaskWithInputFileMonitor):
    def __init__(self, config_stage, config_global):
        super().__init__(config_stage, config_global)
        stage_param = config_stage['parameters']
        self.file_file_name_inputs =  Path(config_global['raw_data_folder']) / stage_param['file_file_name_input']
        self.file_file_name_output =  Path(config_global['raw_data_folder']) / stage_param['file_file_name_output']
        self.file_folder = Path(config_global['file_folder'])
        self.file_types = stage_param['file_types']

    def execute_task(self):

        df_files = pd.read_pickle(self.file_file_name_inputs)

        metadata_list = []
        for index, row in tqdm(df_files.iterrows(), total=df_files.shape[0]):
            metadata_list_sample = {}
            metadata_list_sample['pipe:ID'] = row['pipe:ID']
            metadata_list_sample['pipe:file_type'] = row['pipe:file_type']
            if row['pipe:file_type'] not in self.file_types:
                continue
            if row['pipe:file_type'] not in extractors:
                raise ValueError(f"Extractor for file type '{row['pipe:file_type']}' not found.")

            file_path = self.file_folder / (row['pipe:ID'] + "." + row['pipe:file_type'])
            extractor = extractors[row['pipe:file_type']](file_path, row['pipe:file_type'])
            metadata = extractor.extract()
            if metadata is not None:
                metadata_list_sample.update(metadata)

            metadata_list.append(metadata_list_sample)

        df_metadata_list = pd.DataFrame(metadata_list)
        df_metadata_list['file:created'] = pd.to_datetime(df_metadata_list['file:created'], utc=True)
        df_metadata_list['file:modified'] = pd.to_datetime(df_metadata_list['file:modified'], utc=True)
        df_metadata_list.to_pickle(self.file_file_name_output)

        # Evaluate author names
        nc = NameChecker()
        df_metadata_list.loc[:, 'file:revisedAuthor'] = ""
        for index, row in tqdm(df_metadata_list.iterrows(), total=df_metadata_list.shape[0]):
            if row['file:author'] != "":
                result = nc.get_validated_name(row['file:author'])
                if result != None:
                    df_metadata_list.at[index, 'file:revisedAuthor'] = f"{result.Vorname}/{result.Familienname}"
                    print(f"{result.Vorname} {result.Familienname}")
        df_metadata_list.to_pickle(self.file_file_name_output)